"""
tailor_resume.py
────────────────
Builds the Hyderabad base resume and per-job tailored resumes from the
original .docx, keeping its formatting.

Usage:
    python job-hunt/tailor_resume.py base
    python job-hunt/tailor_resume.py tailor edits.json        # -> private/resume/tailored/
    python job-hunt/tailor_resume.py check  edits.json        # guardrail only

edits.json (written by the agent per job):
{
  "company": "Novartis", "role": "Director Data & Analytics",
  "headline": "Director – Data, Analytics & AI Platforms",
  "keywords": ["Data Strategy", "Power BI", "Data Governance", ...],   # header skills line
  "summary": "Enterprise Data & Analytics leader with 16+ years ...",  # 1st profile paragraph
  "priority_terms": ["governance", "life sciences", "databricks", ...] # used to reorder bullets
}

Guardrail: new text may reword and re-emphasise, but may not introduce a tool,
platform, certification or number that is not already in the base resume.
"""

import json, re, subprocess, sys
from datetime import date
from pathlib import Path

import docx
import yaml

HERE      = Path(__file__).resolve().parent
PRIVATE   = HERE / "private" / "resume"
ORIGINAL  = PRIVATE / "original.docx"
BASE      = PRIVATE / "Lakshmi_Pavani_Base_Hyderabad.docx"
TAILORED  = PRIVATE / "tailored"

PROFILE   = HERE / "private" / "applicant-profile.yaml"

# Specific tools/platforms the guardrail watches for. If one of these appears in
# tailored text but not in the base resume, the edit is rejected.
KNOWN_TOOLS = [
    "AWS", "Redshift", "Glue", "Athena", "EMR", "SageMaker", "Azure Synapse", "Synapse",
    "Azure Data Factory", "ADF", "Purview", "Looker", "Qlik", "QlikView", "Qlik Sense",
    "MicroStrategy", "SAP", "SAP BW", "BusinessObjects", "Spark", "PySpark", "Kafka",
    "dbt", "Airflow", "Hadoop", "Hive", "Teradata", "Netezza", "MongoDB", "Cassandra",
    "Scala", "Java", "R ", "SAS", "SPSS", "TensorFlow", "PyTorch", "Kubernetes", "Docker",
    "Terraform", "Salesforce", "Veeva", "Collibra", "Alation", "Informatica", "Alteryx",
    "Talend", "Matillion", "Fivetran", "Vertex AI", "OpenAI", "Gemini", "LangChain",
    "Power Apps", "SSIS", "SSAS", "SSRS", "Domo", "Sigma", "ThoughtSpot", "Dataiku",
    "PMP", "SAFe", "Prince2", "CSM", "TOGAF", "Six Sigma",
]


# ── helpers ──────────────────────────────────────────────────────────────────

def _para_text(p):
    return "".join(t.text or "" for t in p._p.iter(docx.oxml.ns.qn("w:t")))


def _set_text(p, text):
    """Replace a paragraph's text, keeping the first run's formatting."""
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def _find(doc, startswith):
    for i, p in enumerate(doc.paragraphs):
        if _para_text(p).strip().upper().startswith(startswith.upper()):
            return i
    raise ValueError(f"Section not found: {startswith}")


def _block(doc, start_idx):
    """Contiguous non-empty paragraphs after start_idx (a bullet block)."""
    out, i = [], start_idx + 1
    while i < len(doc.paragraphs) and _para_text(doc.paragraphs[i]).strip():
        out.append(doc.paragraphs[i])
        i += 1
    return out


def _reorder(paras, terms):
    """Stable-sort a block of paragraphs so the ones matching most terms come first."""
    if not paras or not terms:
        return
    terms = [t.lower() for t in terms]
    score = lambda p: -sum(t in _para_text(p).lower() for t in terms)
    ranked = sorted(paras, key=score)
    anchor = paras[0]._p.getprevious()
    for p in ranked:
        anchor.addnext(p._p)
        anchor = p._p


def _reorder_csv(p, terms, sep):
    items = [x.strip() for x in _para_text(p).split(sep.strip()) if x.strip()]
    seen, uniq = set(), []
    for x in items:                     # drop duplicates (e.g. IBM Cognos listed twice)
        if x.lower() not in seen:
            seen.add(x.lower()); uniq.append(x)
    terms = [t.lower() for t in terms or []]
    uniq.sort(key=lambda x: -sum(t in x.lower() or x.lower() in t for t in terms))
    _set_text(p, sep.join(uniq))


def all_text(doc):
    return "\n".join(_para_text(p) for p in doc.paragraphs)


def to_pdf(path):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir",
                    str(path.parent), str(path)], check=True, capture_output=True)
    return path.with_suffix(".pdf")


# ── base resume ──────────────────────────────────────────────────────────────

def build_base():
    # Phone and city come from the private profile so they never land in git.
    me = yaml.safe_load(PROFILE.read_text())["personal"]
    phone, location = me["phone"], f"Location | {me['city']}, {me['country']}"
    doc = docx.Document(ORIGINAL)
    for p in doc.paragraphs[:8]:
        t = _para_text(p)
        if "Mobile" in t:
            for r in p.runs:
                r.text = re.sub(r"\+?\d[\d\s-]{8,}", phone, r.text)
        elif t.strip().lower().startswith("current location"):
            _set_text(p, location)
        elif t.count("|") > 5 and t.rstrip().endswith("|"):
            _set_text(p, t.rstrip().rstrip("|").rstrip())
    for p in doc.paragraphs:
        for r in p.runs:
            r.text = (r.text.replace("IBM Data stage", "IBM DataStage")
                            .replace("Vancouver , Canada", "Vancouver, Canada")
                            .replace(",  ", ", "))
    stack = doc.paragraphs[_find(doc, "TECHNICAL STACK") + 1]
    _reorder_csv(stack, [], " | ")
    doc.save(BASE)
    return BASE


# ── guardrail ────────────────────────────────────────────────────────────────

def check(edits, base_text):
    problems = []
    base_low = base_text.lower()
    new_text = " ".join([edits.get("headline", ""), edits.get("summary", ""),
                         " ".join(edits.get("keywords", []))])
    for tool in KNOWN_TOOLS:
        pat = r"(?<![A-Za-z])" + re.escape(tool.strip()) + r"(?![A-Za-z])"
        if re.search(pat, new_text) and not re.search(pat, base_text):
            problems.append(f"'{tool.strip()}' is not in her resume")
    for num in re.findall(r"\d+\+?%?", new_text):
        if num not in base_text:
            problems.append(f"number '{num}' is not in her resume")
    for kw in edits.get("keywords", []):
        words = [w for w in re.findall(r"[A-Za-z][A-Za-z/&+-]{2,}", kw.lower())]
        if words and not any(w in base_low for w in words):
            problems.append(f"keyword '{kw}' has no basis in her resume")
    return problems


# ── tailoring ────────────────────────────────────────────────────────────────

def tailor(edits):
    if not BASE.exists():
        build_base()
    doc = docx.Document(BASE)
    problems = check(edits, all_text(doc))
    if problems:
        raise SystemExit("Guardrail blocked this tailoring:\n  - " + "\n  - ".join(problems))

    terms = edits.get("priority_terms", [])
    if edits.get("headline"):
        _set_text(doc.paragraphs[1], edits["headline"])
    if edits.get("keywords"):
        _set_text(doc.paragraphs[2], " | ".join(edits["keywords"]))
    if edits.get("summary"):
        _set_text(doc.paragraphs[_find(doc, "EXECUTIVE PROFILE") + 1], edits["summary"])

    _reorder(_block(doc, _find(doc, "Leadership & Business Impact")), terms)
    _reorder(_block(doc, _find(doc, "ENTERPRISE PLATFORM EXPERTISE")), terms)
    _reorder_csv(doc.paragraphs[_find(doc, "TECHNICAL STACK") + 1], terms, " | ")
    for label in ["Enterprise BI & Visualization", "Cloud & Data Platforms",
                  "Engineering & Governance", "AI & Automation"]:
        _reorder_csv(doc.paragraphs[_find(doc, label) + 1], terms, ", ")
    cog = _find(doc, "COGNIZANT")
    _reorder(_block(doc, cog + 1), terms)          # bullets start after the date line

    slug = lambda s: re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-")[:40]
    TAILORED.mkdir(parents=True, exist_ok=True)
    out = TAILORED / f"{date.today():%Y%m%d}_{slug(edits['company'])}_{slug(edits['role'])}.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "base"
    if cmd == "base":
        b = build_base(); print(b); print(to_pdf(b))
    elif cmd in ("tailor", "check"):
        edits = json.loads(Path(sys.argv[2]).read_text())
        if cmd == "check":
            base = docx.Document(BASE if BASE.exists() else build_base())
            probs = check(edits, all_text(base))
            print("OK" if not probs else "\n".join(probs)); sys.exit(1 if probs else 0)
        out = tailor(edits); print(out); print(to_pdf(out))
    else:
        print(__doc__)
