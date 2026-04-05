"""
sheets_manager.py  –  openpyxl-based local .xlsx tracker
No Google Cloud / API key required.
The file is saved locally, pushed to GitHub, and uploaded to Google Drive via browser.
"""

from __future__ import annotations
import os
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, numbers
)
from openpyxl.styles.numbers import FORMAT_NUMBER_COMMA_SEPARATED1
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import (
    TRACKER_FILE, TRACKER_FILE_GIT, SHEET_NAMES, HOUSING_SITES,
    SITES_SHEET_HEADERS, PROPERTY_SHEET_HEADERS, CONTACTS_SHEET_HEADERS,
)

# ── Palette ─────────────────────────────────────────────────────────────────
def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)

def _font(bold=False, color="000000", size=10) -> Font:
    return Font(bold=bold, color=color, size=size, name="Calibri")

def _center() -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def _left() -> Alignment:
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

FILLS = {
    "sites":     _fill("1F2B5E"),   # dark navy
    "property1": _fill("145A32"),   # dark green
    "property2": _fill("1A5276"),   # dark blue
    "contacts":  _fill("6C3483"),   # purple
    "dashboard": _fill("7B241C"),   # dark red
}

WHITE_FONT = _font(bold=True, color="FFFFFF", size=10)
DARK_FONT  = _font(bold=False, color="000000", size=10)
TITLE_FONT = _font(bold=True,  color="1F2B5E", size=13)

LIGHT_FILLS = {
    "sites":     _fill("EAF0FB"),
    "property1": _fill("EAFAF1"),
    "property2": _fill("EBF5FB"),
    "contacts":  _fill("F5EEF8"),
    "dashboard": _fill("FDEDEC"),
}

STATUS_FILLS = {
    "Active":          _fill("D4EFDF"),
    "Failed":          _fill("FADBD8"),
    "Pending":         _fill("FEF9E7"),
    "Manual Required": _fill("FAD7A0"),
    "Expired":         _fill("D5D8DC"),
}

THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)

# ── Low-level helpers ────────────────────────────────────────────────────────

def _set_col_widths(ws: Worksheet, widths: dict[int, int]):
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w

def _header_row(ws: Worksheet, row: int, headers: list[str], fill_key: str):
    fill = FILLS[fill_key]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill = fill
        cell.font = WHITE_FONT
        cell.alignment = _center()
        cell.border = THIN_BORDER

def _title_row(ws: Worksheet, title: str, ncols: int, fill_key: str):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    cell = ws.cell(row=1, column=1, value=title)
    cell.fill = LIGHT_FILLS[fill_key]
    cell.font = TITLE_FONT
    cell.alignment = _center()
    ws.row_dimensions[1].height = 28

def _data_row(ws: Worksheet, row: int, values: list, fill=None):
    for col, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.font = DARK_FONT
        cell.alignment = _left()
        cell.border = THIN_BORDER
        if fill:
            cell.fill = fill

def _freeze(ws: Worksheet, row: int = 2, col: int = 1):
    ws.freeze_panes = ws.cell(row=row + 1, column=col)


# ── Sheet builders ───────────────────────────────────────────────────────────

def _build_sites_sheet(ws: Worksheet):
    ws.title = SHEET_NAMES["sites"]
    ws.sheet_view.showGridLines = True

    _title_row(ws, "Sites & Credentials  –  Fill your Username & Password for each site", len(SITES_SHEET_HEADERS), "sites")
    _header_row(ws, 2, SITES_SHEET_HEADERS, "sites")
    _freeze(ws, 2)

    for i, site in enumerate(HOUSING_SITES, 3):
        vals = [
            site["id"], site["name"], site["url"], site["post_url"],
            site["tier"], "Yes" if site["free_listing"] else "No",
            "", "", "", "", "Not Set", "",
        ]
        _data_row(ws, i, vals)
        # Zebra stripes
        if i % 2 == 0:
            for col in range(1, len(SITES_SHEET_HEADERS) + 1):
                ws.cell(i, col).fill = _fill("F2F3F4")

    _set_col_widths(ws, {1:12, 2:18, 3:28, 4:32, 5:10, 6:12, 7:24, 8:18, 9:18, 10:18, 11:14, 12:26})
    ws.row_dimensions[2].height = 20


def _build_property_sheet(ws: Worksheet, sheet_key: str, prop_title: str):
    ws.title = SHEET_NAMES[sheet_key]
    ncols = len(PROPERTY_SHEET_HEADERS)

    _title_row(ws, f"Ad Status Tracker  –  {prop_title}", ncols, sheet_key)
    _header_row(ws, 2, PROPERTY_SHEET_HEADERS, sheet_key)
    _freeze(ws, 2)

    for i, site in enumerate(HOUSING_SITES, 3):
        vals = [site["id"], site["name"], "", "", "", "Pending", 0, 0, 0, "", "", "", ""]
        _data_row(ws, i, vals, fill=STATUS_FILLS["Pending"])

    _set_col_widths(ws, {1:12, 2:18, 3:14, 4:36, 5:36, 6:14, 7:8, 8:8, 9:10, 10:16, 11:14, 12:14, 13:30})
    ws.row_dimensions[2].height = 20

    # Add data-validation-style notes for Status column (col F = 6)
    ws.cell(2, 6).comment = None  # placeholder; openpyxl comments need extra import if needed


def _build_contacts_sheet(ws: Worksheet):
    ws.title = SHEET_NAMES["contacts"]
    ncols = len(CONTACTS_SHEET_HEADERS)

    _title_row(ws, "Contact Tracker  –  Update Follow-up Status after every call/message", ncols, "contacts")
    _header_row(ws, 2, CONTACTS_SHEET_HEADERS, "contacts")
    _freeze(ws, 2)

    _set_col_widths(ws, {1:12, 2:8, 3:14, 4:18, 5:20, 6:16, 7:24, 8:34, 9:18, 10:16, 11:16, 12:14, 13:14, 14:30})
    ws.row_dimensions[2].height = 20


def _build_dashboard_sheet(ws: Worksheet, p1_name: str, p2_name: str):
    ws.title = SHEET_NAMES["dashboard"]
    ws.sheet_view.showGridLines = False

    # ── Title ──────────────────────────────────────────────────────────────
    ws.merge_cells("A1:F1")
    t = ws.cell(1, 1, "HOUSING ADS DASHBOARD  –  Hyderabad")
    t.font = _font(bold=True, color="FFFFFF", size=14)
    t.fill = FILLS["dashboard"]
    t.alignment = _center()
    ws.row_dimensions[1].height = 32

    ws.cell(2, 1, "Last Refreshed:").font = _font(bold=True)
    ws.cell(2, 2, f'=TEXT(NOW(),"DD-MM-YYYY HH:MM")')
    ws.cell(2, 2).font = _font(color="7B241C")

    # ── Metrics table ──────────────────────────────────────────────────────
    p1 = SHEET_NAMES["property1"]
    p2 = SHEET_NAMES["property2"]
    con = SHEET_NAMES["contacts"]

    def hdr(row, col, val):
        c = ws.cell(row, col, val)
        c.font = _font(bold=True, color="FFFFFF")
        c.fill = FILLS["dashboard"]
        c.alignment = _center()
        c.border = THIN_BORDER

    def metric(row, col, val):
        c = ws.cell(row, col, val)
        c.font = _font(size=11)
        c.alignment = _center()
        c.border = THIN_BORDER

    # Section header
    ws.merge_cells("A4:F4")
    s = ws.cell(4, 1, "AD STATUS OVERVIEW")
    s.font = _font(bold=True, color="FFFFFF", size=11)
    s.fill = _fill("2E4057")
    s.alignment = _center()

    hdr(5, 1, "Metric"); hdr(5, 2, p1_name); hdr(5, 3, p2_name); hdr(5, 4, "TOTAL")

    metrics = [
        ("Ads – Active",   f"=COUNTIF('{p1}'!F3:F100,\"Active\")",   f"=COUNTIF('{p2}'!F3:F100,\"Active\")",   "=B{r}+C{r}"),
        ("Ads – Pending",  f"=COUNTIF('{p1}'!F3:F100,\"Pending\")",  f"=COUNTIF('{p2}'!F3:F100,\"Pending\")",  "=B{r}+C{r}"),
        ("Ads – Failed",   f"=COUNTIF('{p1}'!F3:F100,\"Failed\")",   f"=COUNTIF('{p2}'!F3:F100,\"Failed\")",   "=B{r}+C{r}"),
        ("Total Views",    f"=SUM('{p1}'!G3:G100)",                  f"=SUM('{p2}'!G3:G100)",                  "=B{r}+C{r}"),
        ("Total Clicks",   f"=SUM('{p1}'!H3:H100)",                  f"=SUM('{p2}'!H3:H100)",                  "=B{r}+C{r}"),
        ("Total Inquiries",f"=SUM('{p1}'!I3:I100)",                  f"=SUM('{p2}'!I3:I100)",                  "=B{r}+C{r}"),
    ]
    for i, (label, f1, f2, ftot) in enumerate(metrics, 6):
        metric(i, 1, label); metric(i, 2, f1); metric(i, 3, f2)
        metric(i, 4, ftot.format(r=i))
        ws.cell(i, 1).font = _font(bold=True)

    # ── Contact pipeline ──────────────────────────────────────────────────
    ws.merge_cells("A13:F13")
    s2 = ws.cell(13, 1, "BUYER CONTACT PIPELINE")
    s2.font = _font(bold=True, color="FFFFFF", size=11)
    s2.fill = _fill("6C3483")
    s2.alignment = _center()

    pipeline = [
        ("New",             f"=COUNTIF('{con}'!I3:I200,\"New\")"),
        ("Interested",      f"=COUNTIF('{con}'!I3:I200,\"Interested\")"),
        ("Very Interested",f"=COUNTIF('{con}'!I3:I200,\"Very Interested\")"),
        ("Negotiating",     f"=COUNTIF('{con}'!I3:I200,\"Negotiating\")"),
        ("Closed / Sold",   f"=COUNTIF('{con}'!I3:I200,\"Closed / Sold\")"),
        ("Not Interested",  f"=COUNTIF('{con}'!I3:I200,\"Not Interested\")"),
        ("No Response",     f"=COUNTIF('{con}'!I3:I200,\"No Response\")"),
    ]
    PIPELINE_FILLS = {
        "New": _fill("D6EAF8"), "Interested": _fill("FEF9E7"),
        "Very Interested": _fill("D4EFDF"), "Negotiating": _fill("D4EFDF"),
        "Closed / Sold": _fill("A9DFBF"), "Not Interested": _fill("FADBD8"),
        "No Response": _fill("EAECEE"),
    }
    for i, (label, formula) in enumerate(pipeline, 14):
        c1 = ws.cell(i, 1, label)
        c1.font = _font(bold=True)
        c1.fill = PIPELINE_FILLS.get(label, _fill("FFFFFF"))
        c1.border = THIN_BORDER
        c1.alignment = _left()
        c2 = ws.cell(i, 2, formula)
        c2.font = _font(bold=True, size=12)
        c2.fill = PIPELINE_FILLS.get(label, _fill("FFFFFF"))
        c2.border = THIN_BORDER
        c2.alignment = _center()

    # ── Trends header ─────────────────────────────────────────────────────
    ws.merge_cells("A22:F22")
    s3 = ws.cell(22, 1, "TOP SITES BY PERFORMANCE  (refresh via: python analytics.py)")
    s3.font = _font(bold=True, color="FFFFFF", size=11)
    s3.fill = _fill("1F2B5E")
    s3.alignment = _center()

    _set_col_widths(ws, {1:24, 2:18, 3:18, 4:12, 5:12, 6:12})
    ws.freeze_panes = "A2"


# ── Public API ───────────────────────────────────────────────────────────────

def create_tracker(prop1_title: str = "Property 1 – Hyderabad",
                   prop2_title: str = "Property 2 – Hyderabad") -> Path:
    """
    Create (or recreate) the full .xlsx tracker file.
    Returns the path to the saved file.
    """
    TRACKER_FILE.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Add sheets in order
    _build_sites_sheet(wb.create_sheet(SHEET_NAMES["sites"]))
    _build_property_sheet(wb.create_sheet(SHEET_NAMES["property1"]), "property1", prop1_title)
    _build_property_sheet(wb.create_sheet(SHEET_NAMES["property2"]), "property2", prop2_title)
    _build_contacts_sheet(wb.create_sheet(SHEET_NAMES["contacts"]))
    _build_dashboard_sheet(wb.create_sheet(SHEET_NAMES["dashboard"]), prop1_title, prop2_title)

    # Save to user's preferred location
    try:
        TRACKER_FILE.parent.mkdir(parents=True, exist_ok=True)
        wb.save(TRACKER_FILE)
    except Exception:
        pass   # Windows path not accessible from Linux dev env – that's fine

    # Always save a git-trackable copy inside the project dir
    if TRACKER_FILE_GIT != TRACKER_FILE:
        TRACKER_FILE_GIT.parent.mkdir(parents=True, exist_ok=True)
        wb.save(TRACKER_FILE_GIT)

    return TRACKER_FILE_GIT if TRACKER_FILE_GIT.exists() else TRACKER_FILE


def _load_wb():
    # Prefer the git copy; fall back to the user-path copy
    target = TRACKER_FILE_GIT if TRACKER_FILE_GIT.exists() else TRACKER_FILE
    if not target.exists():
        raise FileNotFoundError(
            f"Tracker file not found: {TRACKER_FILE}\n"
            "Run setup_sheets.py first."
        )
    return load_workbook(target)


def _save_wb(wb: Workbook):
    # Save to git copy (always accessible)
    TRACKER_FILE_GIT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(TRACKER_FILE_GIT)
    # Also save to user's preferred path if accessible
    if TRACKER_FILE != TRACKER_FILE_GIT:
        try:
            TRACKER_FILE.parent.mkdir(parents=True, exist_ok=True)
            wb.save(TRACKER_FILE)
        except Exception:
            pass


def get_credentials() -> dict[str, dict]:
    """Return {site_id: {username, password, phone, api_key}}."""
    wb = _load_wb()
    ws = wb[SHEET_NAMES["sites"]]
    creds = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or not row[0]:
            continue
        site_id, _, _, _, _, _, username, password, phone, api_key = (list(row) + [""] * 10)[:10]
        if username:
            creds[str(site_id)] = {
                "username": str(username),
                "password": str(password) if password else "",
                "phone":    str(phone)    if phone    else "",
                "api_key":  str(api_key)  if api_key  else "",
            }
    return creds


def update_ad_status(property_key: str, site_id: str, updates: dict):
    """Update a row in the property status sheet and re-save the file."""
    wb = _load_wb()
    ws = wb[SHEET_NAMES[property_key]]

    STATUS_COL   = 6
    VIEWS_COL    = 7
    CLICKS_COL   = 8
    INQ_COL      = 9
    DATE_COL     = 3
    URL_COL      = 4
    TITLE_COL    = 5
    SCRAPED_COL  = 10
    EXPIRY_COL   = 11
    NOTES_COL    = 13

    for row in ws.iter_rows(min_row=3):
        if row[0].value == site_id:
            r = row[0].row
            def _set(col, key):
                if key in updates:
                    ws.cell(r, col).value = updates[key]
                    ws.cell(r, col).font = DARK_FONT
                    ws.cell(r, col).alignment = _left()
                    ws.cell(r, col).border = THIN_BORDER

            _set(DATE_COL,  "posted_date")
            _set(URL_COL,   "ad_url")
            _set(TITLE_COL, "ad_title")
            _set(VIEWS_COL, "views")
            _set(CLICKS_COL,"clicks")
            _set(INQ_COL,   "inquiries")
            _set(SCRAPED_COL,"last_scraped")
            _set(EXPIRY_COL,"expiry_date")
            _set(NOTES_COL, "notes")

            if "status" in updates:
                status = updates["status"]
                cell = ws.cell(r, STATUS_COL)
                cell.value = status
                cell.font = _font(bold=True)
                cell.alignment = _center()
                cell.border = THIN_BORDER
                cell.fill = STATUS_FILLS.get(status, _fill("FFFFFF"))
            break

    _save_wb(wb)


def add_contact(contact: dict):
    """Append a new contact row to the Contacts sheet."""
    wb = _load_wb()
    ws = wb[SHEET_NAMES["contacts"]]
    now = datetime.now()

    # Find next empty row
    next_row = ws.max_row + 1
    if ws.max_row <= 2:
        next_row = 3

    values = [
        now.strftime("%d-%m-%Y"),
        now.strftime("%H:%M"),
        contact.get("property", ""),
        contact.get("site", ""),
        contact.get("name", ""),
        contact.get("phone", ""),
        contact.get("email", ""),
        contact.get("message", ""),
        "New",
        "", "", "", "",
        contact.get("notes", ""),
    ]
    _data_row(ws, next_row, values)
    _save_wb(wb)


def update_contact_row(row_num: int, updates: dict):
    """Update specific columns in a contact row (1-indexed)."""
    wb = _load_wb()
    ws = wb[SHEET_NAMES["contacts"]]

    mapping = {
        "follow_up_status":  9,
        "last_follow_up":   10,
        "next_follow_up":   11,
        "asking_price":     12,
        "offered_price":    13,
        "notes":            14,
    }
    for key, col in mapping.items():
        if key in updates:
            ws.cell(row_num, col).value = updates[key]
            ws.cell(row_num, col).font = DARK_FONT
            ws.cell(row_num, col).alignment = _left()
            ws.cell(row_num, col).border = THIN_BORDER
    _save_wb(wb)


def load_contacts() -> list[dict]:
    """Return all contacts from the Contacts sheet."""
    wb = _load_wb()
    ws = wb[SHEET_NAMES["contacts"]]
    contacts = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or not row[4]:   # need a name
            continue
        contacts.append({
            "row":              ws.min_row + contacts.__len__() + 2,
            "date":             str(row[0]  or ""),
            "time":             str(row[1]  or ""),
            "property":         str(row[2]  or ""),
            "site":             str(row[3]  or ""),
            "name":             str(row[4]  or ""),
            "phone":            str(row[5]  or ""),
            "email":            str(row[6]  or ""),
            "message":          str(row[7]  or ""),
            "follow_up_status": str(row[8]  or "New"),
            "last_follow_up":   str(row[9]  or ""),
            "next_follow_up":   str(row[10] or ""),
            "asking_price":     str(row[11] or ""),
            "offered_price":    str(row[12] or ""),
            "notes":            str(row[13] or ""),
        })
    return contacts


def load_property_stats(property_key: str) -> list[dict]:
    """Return all rows from a property status sheet."""
    wb = _load_wb()
    ws = wb[SHEET_NAMES[property_key]]
    rows = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or not row[0]:
            continue
        rows.append({
            "site_id":    str(row[0] or ""),
            "site_name":  str(row[1] or ""),
            "posted_date":str(row[2] or ""),
            "ad_url":     str(row[3] or ""),
            "ad_title":   str(row[4] or ""),
            "status":     str(row[5] or "Pending"),
            "views":      int(row[6] or 0),
            "clicks":     int(row[7] or 0),
            "inquiries":  int(row[8] or 0),
            "last_scraped":str(row[9] or ""),
        })
    return rows


def write_trends_table(trends: list[dict]):
    """Write per-site trend data into the Dashboard sheet starting at row 23."""
    wb = _load_wb()
    ws = wb[SHEET_NAMES["dashboard"]]

    START = 23
    headers = ["Site", "P1 Views", "P1 Clicks", "P1 Inquiries",
               "P2 Views", "P2 Clicks", "P2 Inquiries", "Total Views"]
    _header_row(ws, START, headers, "dashboard")

    for i, t in enumerate(trends, START + 1):
        vals = [t["name"], t["p1_views"], t["p1_clicks"], t["p1_inq"],
                t["p2_views"], t["p2_clicks"], t["p2_inq"],
                t["p1_views"] + t["p2_views"]]
        _data_row(ws, i, vals)
        if (i - START) % 2 == 0:
            for col in range(1, 9):
                ws.cell(i, col).fill = _fill("F2F3F4")

    # Bold the total column
    for i in range(START + 1, START + 1 + len(trends)):
        ws.cell(i, 8).font = _font(bold=True, size=11)

    _save_wb(wb)
